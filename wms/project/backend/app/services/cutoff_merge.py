# 截单文件合并边界：仅处理出库单工作表，按文本出库单号对齐，不重新分析波次。
"""Merge the WMS order sheet only; one outbound order is one complete row."""
from __future__ import annotations

import os
import re
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell


def normalized(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


REQUIRED = {"outboundorderno/出库单号", "shippingcarrier/物流承运商",
            "typeofordervariety/订单品种类型", "package1trackingno./物流跟踪号",
            "sku1sku", "totalqtyofsku/总数量"}


def read_orders(path: Path) -> tuple[list[str], dict[str, tuple]]:
    # 截单合并采用严格的一行一单契约：同文件重复单号直接报错，不能猜测如何合并。
    # 保留文本单号和整行原值；公式不是确定的订单原始值，拒绝使用缓存结果代替。
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if "出库单" not in book.sheetnames:
            raise ValueError(f"{path.name} 缺少出库单工作表")
        sheet = book["出库单"]
        rows = sheet.iter_rows()
        header = [str(c.value or "").strip() for c in next(rows, [])]
        keys = [normalized(h) for h in header]
        if not REQUIRED.issubset(keys) or len(set(keys)) != len(keys) or "" in keys:
            raise ValueError(f"{path.name} 表头缺失或重复")
        index = keys.index("outboundorderno/出库单号")
        orders = {}
        for row_number, cells in enumerate(rows, start=2):
            values = tuple(None if c.value == "" else c.value for c in cells)
            values += (None,) * (len(header) - len(values))
            if all(v is None or v == "" for v in values):
                continue
            if any(c.data_type == "f" for c in cells):
                raise ValueError(f"{path.name} 含公式，无法确认订单原始值")
            raw = values[index]
            if not isinstance(raw, str) or not raw.strip():
                raise ValueError(f"{path.name} 第 {row_number} 行缺少文本出库单号")
            key = raw.strip()
            if key in orders:
                raise ValueError(f"{path.name} 内出库单号重复，不符合一行一单约定")
            orders[key] = values
        return header, orders
    finally:
        book.close()


def merge_orders(base: Path, delta: Path, target: Path) -> dict[str, int]:
    # 两次导出的动态 SKU 列数可能不同，先按归一化列名对齐，再按订单号取并集。
    # 同号以增量整行覆盖，不逐格补值；否则旧 SKU 可能残留在已经变化的订单中。
    # 结果先写临时文件，完整回读比对后再发布 target，失败不留下“成功”的半成品。
    headers, old = read_orders(base)
    extra_headers, new = read_orders(delta)
    keys, extra = list(map(normalized, headers)), list(map(normalized, extra_headers))
    base_keys = keys.copy()
    # WMS expands SKU n columns according to the largest order in each export.
    fixed = lambda seq: {k for k in seq if not re.match(r"^sku\d+", k)}
    if fixed(keys) != fixed(extra):
        raise ValueError("两次导出的模板字段不一致，未发布合并文件")
    for name, key in zip(extra_headers, extra):
        if key not in keys:
            headers.append(name)
            keys.append(key)
    def align(values, source):
        mapping = dict(zip(source, values))
        return tuple(mapping.get(k) for k in keys)
    merged = {key: align(row, base_keys) for key, row in old.items()}
    # Base column positions are retained; newly appended columns are blank.
    for key, row in new.items():
        merged[key] = align(row, extra)
    overlap = len(old.keys() & new.keys())
    if len(merged) != len(old) + len(new) - overlap:
        raise ValueError("合并订单数校验失败")
    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_name(f".{target.stem}-{uuid4().hex}.xlsx")
    book = Workbook(write_only=True)
    sheet = book.create_sheet("出库单")
    try:
        sheet.append(headers)
        for values in merged.values():
            cells = []
            for value in values:
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cells.append(cell)
            sheet.append(cells)
        book.save(tmp)
        _, verified = read_orders(tmp)
        if verified != merged:
            raise ValueError("合并文件回读校验失败")
        os.replace(tmp, target)
    finally:
        book.close()
        tmp.unlink(missing_ok=True)
    return {"base": len(old), "delta": len(new), "overlap": overlap, "total": len(merged)}
