"""双表订单关联、SKU 对齐、原子发布和截单库位保留；仅使用临时工作簿。"""
import asyncio
import threading
from contextlib import asynccontextmanager
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import pytest
from openpyxl import Workbook, load_workbook

from app.services.order_workbook import normalize_order_file, read_orders, normalized
from app.services.cutoff_merge import merge_orders

HEAD = ['Outbound Order No/出库单号', 'Shipping Carrier/物流承运商',
        'Type of order variety/订单品种类型', 'Package 1 Tracking No./物流跟踪号',
        'Total Qty of SKU/总数量', 'Client/客户', 'SKU 1 SKU', 'SKU 2 SKU']
DETAIL = ['Outbound Order No/出库单号', 'SKU 1 SKU', 'SKU 1 Outbound Qty/出库数量',
          'SKU 1 Recommended Picking Location/推荐拣货库位1',
          'SKU 2 SKU', 'SKU 2 Outbound Qty/出库数量',
          'SKU 2 Recommended Picking Location/推荐拣货库位1']


def fixture(path, main=None, details=None):
    book = Workbook()
    sheet = book.active
    sheet.title = '出库单'
    sheet.append(HEAD)
    for row in main if main is not None else [
        ['000A', 'USPS', '多品混合', '001234', 3, '客户(C)', '001SKU', 'B'],
        ['000B', 'CBT', '单品单件', '001235', 1, '客户(C)', 'C', None],
    ]: sheet.append(row)
    sheet = book.create_sheet('拣货库位明细')
    sheet.append(DETAIL)
    # 订单顺序与 SKU 槽位均与主表相反，确保用身份关联而非行号/列号。
    for row in details if details is not None else [
        ['000B', 'C', 1, 'C-01', None, None, None],
        ['000A', 'B', 2, 'B-01', '001SKU', 1, 'A-01'],
    ]: sheet.append(row)
    book.save(path)
    book.close()
    return path


def test_normalize_reorders_by_order_and_sku_preserves_values_and_is_idempotent(tmp_path):
    path = fixture(tmp_path/'orders.xlsx')
    assert normalize_order_file(path) == 2
    book = load_workbook(path, read_only=True)
    assert book.sheetnames == ['出库单']
    book.close()
    headers, orders = read_orders(path)
    assert len(headers) == len(HEAD)+4
    assert len(set(map(normalized, headers))) == len(headers)
    row = dict(zip(map(normalized, headers), orders['000A']))
    assert row['sku1sku'] == '001SKU'
    assert row['sku1outboundqty/出库数量'] == 1
    assert row['sku1recommendedpickinglocation/推荐拣货库位1'] == 'A-01'
    assert row['sku2outboundqty/出库数量'] == 2
    assert row['sku2recommendedpickinglocation/推荐拣货库位1'] == 'B-01'
    assert row['package1trackingno./物流跟踪号'] == '001234'
    assert normalize_order_file(path) == 2
    assert read_orders(path) == (headers, orders)


@pytest.mark.parametrize('case', ['missing', 'extra', 'duplicate', 'sku', 'qty', 'formula', 'blank_sku_data'])
def test_invalid_detail_preserves_original_download(tmp_path, case):
    path = fixture(tmp_path/'bad.xlsx')
    book = load_workbook(path)
    detail = book['拣货库位明细']
    if case == 'missing': detail.delete_rows(2)
    if case == 'extra': detail.append(['EXTRA', 'C', 1, 'X'])
    if case == 'duplicate': detail.append(['000A', 'B', 2, 'X'])
    if case == 'sku': detail['B3'] = 'WRONG'
    if case == 'qty': detail['C3'] = 999
    if case == 'formula': detail['C3'] = '=1+1'
    if case == 'blank_sku_data': detail['G2'] = 'ORPHAN'
    book.save(path)
    book.close()
    before = path.read_bytes()
    with pytest.raises(ValueError): normalize_order_file(path)
    assert path.read_bytes() == before
    assert not list(tmp_path.glob('.*.xlsx'))


def test_multiple_location_columns_and_empty_workbook(tmp_path):
    path = fixture(tmp_path/'multi.xlsx')
    book = load_workbook(path)
    sheet = book['拣货库位明细']
    sheet.cell(1, 8, 'SKU 1 Recommended Picking Location/推荐拣货库位2')
    sheet.cell(3, 8, 'B-02')
    book.save(path);book.close()
    normalize_order_file(path)
    headers, orders = read_orders(path)
    row = dict(zip(map(normalized, headers), orders['000A']))
    assert row['sku2recommendedpickinglocation/推荐拣货库位2'] == 'B-02'
    assert row['sku1recommendedpickinglocation/推荐拣货库位2'] is None
    empty = fixture(tmp_path/'empty.xlsx', [], [])
    assert normalize_order_file(empty) == 0


def test_cutoff_joins_raw_sheets_and_replaces_all_old_sku_locations(tmp_path):
    base = fixture(tmp_path/'base.xlsx')
    delta = fixture(tmp_path/'delta.xlsx',
        [['000A', 'USPS', '单品单件', '001234', 1, '客户(C)', 'NEW', None]],
        [['000A', 'NEW', 1, 'NEW-LOC', None, None, None]])
    target = tmp_path/'cutoff.xlsx'
    assert merge_orders(base, delta, target) == dict(base=2, delta=1, overlap=1, total=2)
    headers, orders = read_orders(target)
    row = dict(zip(map(normalized, headers), orders['000A']))
    assert row['sku1sku'] == 'NEW'
    assert row['sku1recommendedpickinglocation/推荐拣货库位1'] == 'NEW-LOC'
    assert row['sku2sku'] is None
    assert row['sku2recommendedpickinglocation/推荐拣货库位1'] is None


def test_failed_atomic_publication_keeps_original(tmp_path, monkeypatch):
    path = fixture(tmp_path/'orders.xlsx');before = path.read_bytes()
    def fail(*_): raise OSError('replace failed')
    monkeypatch.setattr('app.services.order_workbook.os.replace', fail)
    with pytest.raises(OSError): normalize_order_file(path)
    assert path.read_bytes() == before
    assert not list(tmp_path.glob('.*.xlsx'))


def test_missing_detail_and_mixed_legacy_cutoff_fail_without_losing_source(tmp_path):
    legacy = fixture(tmp_path/'legacy.xlsx')
    book = load_workbook(legacy);del book['拣货库位明细'];book.save(legacy);book.close()
    before = legacy.read_bytes()
    with pytest.raises(ValueError, match='缺少拣货库位'): normalize_order_file(legacy)
    assert legacy.read_bytes() == before
    assert len(read_orders(legacy)[1]) == 2  # 旧文件仍可读取，不强制改写历史数据。
    delta = fixture(tmp_path/'delta.xlsx')
    with pytest.raises(ValueError, match='库位字段不一致'):
        merge_orders(legacy, delta, tmp_path/'result.xlsx')
    assert not (tmp_path/'result.xlsx').exists()


@pytest.mark.asyncio
@pytest.mark.parametrize('scheduled', [False, True])
@pytest.mark.parametrize('invalid', [False, True])
async def test_export_only_delivers_normalized_workbook(isolated_runtime, monkeypatch, scheduled, invalid):
    from app.automation import wms_export as mod
    path = fixture(isolated_runtime.downloads_dir/'orders.xlsx')
    if invalid:
        book = load_workbook(path);book['拣货库位明细']['C3'] = 999;book.save(path);book.close()
    before = path.read_bytes()
    @asynccontextmanager
    async def browser(*_, **__): yield SimpleNamespace(pages=[SimpleNamespace(url='about:blank')])
    monkeypatch.setattr(mod, 'open_browser_context', browser)
    automation = mod.WmsExportAutomation(isolated_runtime, isolated_runtime.load_automation())
    automation._prepare_and_submit = AsyncMock(return_value=('done', '渠道拆分', SimpleNamespace(filename=path.name)))
    automation.task_center.download = AsyncMock(return_value=path)
    progress = AsyncMock()
    if invalid:
        with pytest.raises(mod.AutomationError, match='原始下载已保留'):
            await automation.run(progress, headless=True, created_range={'start':'2026-10-01', 'end':'2026-10-02'} if scheduled else None)
        assert path.read_bytes() == before
        assert not any(c.args[0] == 'orders_normalized' for c in progress.call_args_list)
        return
    result = await automation.run(progress, headless=True, created_range={'start':'2026-10-01', 'end':'2026-10-02'} if scheduled else None)
    assert result.downloaded_file == str(path)
    book = load_workbook(path, read_only=True)
    assert book.sheetnames == ['出库单'];book.close()
    assert any(c.args[0] == 'orders_normalized' for c in progress.call_args_list)


@pytest.mark.asyncio
async def test_cancel_waits_for_local_normalization_before_releasing_context(isolated_runtime, monkeypatch):
    from app.automation import wms_export as mod
    started, release = threading.Event(), threading.Event()
    closed = []
    def normalize(_):
        started.set()
        assert release.wait(5)
        return 1
    @asynccontextmanager
    async def browser(*_, **__):
        try: yield SimpleNamespace(pages=[SimpleNamespace(url='about:blank')])
        finally: closed.append(True)
    monkeypatch.setattr(mod, 'open_browser_context', browser)
    monkeypatch.setattr(mod, 'normalize_order_file', normalize)
    automation = mod.WmsExportAutomation(isolated_runtime, isolated_runtime.load_automation())
    automation._prepare_and_submit = AsyncMock(return_value=('done', '渠道拆分', SimpleNamespace(filename='orders.xlsx')))
    automation.task_center.download = AsyncMock(return_value=isolated_runtime.downloads_dir/'orders.xlsx')
    task = asyncio.create_task(automation.run(AsyncMock(), headless=True))
    try:
        assert await asyncio.to_thread(started.wait, 2)
        task.cancel()
        await asyncio.sleep(0)
        assert not task.done() and not closed
    finally:
        release.set()
        with pytest.raises(asyncio.CancelledError): await task
    assert closed == [True]
