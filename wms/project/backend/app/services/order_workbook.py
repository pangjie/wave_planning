"""订单工作簿边界：按订单号、SKU 合并库位明细，校验后原子发布单表文件。"""
from __future__ import annotations

import math
import os
import re
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell


def normalized(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


ORDER = "outboundorderno/出库单号"
REQUIRED = {ORDER, "shippingcarrier/物流承运商", "typeofordervariety/订单品种类型",
            "package1trackingno./物流跟踪号", "sku1sku", "totalqtyofsku/总数量"}
SKU = re.compile(r"^sku(\d+)sku$")
FIELD = re.compile(r"^sku(\d+)(.+)$")
QTY = "outboundqty/出库数量"
LOCATION = "recommendedpickinglocation/推荐拣货库位"


def _read_sheet(sheet) -> tuple[list[str], dict[str, tuple]]:
    rows = sheet.iter_rows()
    headers = [str(c.value or "").strip() for c in next(rows, [])]
    keys = list(map(normalized, headers))
    if ORDER not in keys or "" in keys or len(set(keys)) != len(keys):
        raise ValueError(f"{sheet.title} 表头缺失或重复")
    index = keys.index(ORDER)
    orders = {}
    for number, cells in enumerate(rows, 2):
        values = tuple(None if c.value == "" else c.value for c in cells)
        # write-only XLSX 会省略行末空单元格；恢复到表头宽度以保证逐列校验一致。
        values += (None,) * (len(headers) - len(values))
        if all(v is None for v in values):
            continue
        if any(c.data_type == "f" for c in cells):
            raise ValueError(f"{sheet.title} 含公式，无法确认订单原始值")
        raw = values[index]
        if not isinstance(raw, str) or not raw.strip():
            raise ValueError(f"{sheet.title} 第 {number} 行缺少文本出库单号")
        key = raw.strip()
        if key in orders:
            raise ValueError(f"{sheet.title} 内出库单号重复，不符合一行一单约定")
        orders[key] = values
    return headers, orders


def _quantity(value, label: str) -> int:
    try:
        n = float(value)
    except (ValueError, TypeError):
        raise ValueError(f"{label} 数量无效") from None
    if isinstance(value, bool) or not math.isfinite(n) or n < 0 or not n.is_integer():
        raise ValueError(f"{label} 数量无效")
    return int(n)


def _sku_columns(keys):
    return {int(m[1]): i for i, k in enumerate(keys) if (m := SKU.fullmatch(k))}


def _sku_values(row, columns, label):
    result = {}
    for slot, index in columns.items():
        value = row[index]
        if value is None:
            continue
        if not isinstance(value, str) or not value.strip():
            raise ValueError(f"{label} SKU 必须是非空文本")
        sku = value.strip()
        if sku in result:
            raise ValueError(f"{label} SKU 重复，无法唯一对应库位")
        result[sku] = slot
    return result


def _join_details(headers, orders, detail_headers, details):
    if orders.keys() != details.keys():
        raise ValueError(f"库位明细订单不一致：缺少 {len(orders.keys()-details.keys())} 单，额外 {len(details.keys()-orders.keys())} 单")
    keys = list(map(normalized, headers))
    detail_keys = list(map(normalized, detail_headers))
    main_slots, detail_slots = _sku_columns(keys), _sku_columns(detail_keys)
    fields, labels = {}, {}
    for index, key in enumerate(detail_keys):
        if key == ORDER or SKU.fullmatch(key):
            continue  # 订单号及 SKU 已在主表保留，不再次输出。
        match = FIELD.fullmatch(key)
        if not match or int(match[1]) not in detail_slots:
            raise ValueError(f"未支持的库位明细列：{detail_headers[index]}")
        slot, suffix = int(match[1]), match[2]
        fields.setdefault(slot, {})[suffix] = index
        labels.setdefault(suffix, re.sub(r"^SKU\s*\d+\s*", "", detail_headers[index], flags=re.I))
    for slot in detail_slots:
        if QTY not in fields.get(slot, {}) or not any(k.startswith(LOCATION) for k in fields[slot]):
            raise ValueError(f"库位明细 SKU {slot} 缺少出库数量或推荐拣货库位列")
    # 原 SKU 编号是列的位置，不是关联键：明细中 SKU 排序不同也必须按真实 SKU 对齐。
    output_headers, mapping = [], []
    for index, (name, key) in enumerate(zip(headers, keys)):
        output_headers.append(name)
        mapping.append((index, None, None))
        match = SKU.fullmatch(key)
        if match:
            for suffix, label in labels.items():
                if f"sku{match[1]}{suffix}" in keys:
                    raise ValueError("出库单已含库位字段却仍有明细表，无法确定数据来源")
                output_headers.append(f"SKU {match[1]}\n{label}")
                mapping.append((None, int(match[1]), suffix))
    total_index = keys.index("totalqtyofsku/总数量")
    joined = {}
    for order, row in orders.items():
        detail = details[order]
        main_skus = _sku_values(row, main_slots, order)
        detail_skus = _sku_values(detail, detail_slots, order)
        if not main_skus or main_skus.keys() != detail_skus.keys():
            raise ValueError(f"订单 {order} 的 SKU 与库位明细不一致")
        total = 0
        for sku, slot in detail_skus.items():
            total += _quantity(detail[fields[slot][QTY]], f"订单 {order} SKU {sku}")
        if total != _quantity(row[total_index], f"订单 {order}"):
            raise ValueError(f"订单 {order} 的 SKU 出库数量之和与总数量不一致")
        for slot in detail_slots.keys()-set(detail_skus.values()):
            if any(detail[i] not in (None, 0) for i in fields[slot].values()):
                raise ValueError(f"订单 {order} 空 SKU 下存在库位或数量数据")
        reverse_slots = {slot: sku for sku, slot in main_skus.items()}
        values = []
        for original, slot, suffix in mapping:
            if original is not None:
                values.append(row[original])
            elif slot not in reverse_slots:
                values.append(None)
            else:
                detail_slot = detail_skus[reverse_slots[slot]]
                col = fields[detail_slot].get(suffix)
                values.append(detail[col] if col is not None else None)
        joined[order] = tuple(values)
    return output_headers, joined


def read_orders(path: Path, *, require_locations: bool = False) -> tuple[list[str], dict[str, tuple]]:
    """兼容旧单表与新双表；新导出必须有明细或已合并字段，不能静默遗漏库位。"""
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if "出库单" not in book.sheetnames:
            raise ValueError(f"{path.name} 缺少出库单工作表")
        headers, orders = _read_sheet(book["出库单"])
        keys = list(map(normalized, headers))
        if not REQUIRED.issubset(keys):
            raise ValueError(f"{path.name} 表头缺失或重复")
        if "拣货库位明细" in book.sheetnames:
            return _join_details(headers, orders, *_read_sheet(book["拣货库位明细"]))
        if require_locations and not any(LOCATION in k for k in keys):
            raise ValueError("导出文件缺少拣货库位明细，请确认渠道拆分模板已勾选推荐拣货库位")
        return headers, orders
    finally:
        book.close()


def write_orders(target: Path, headers: list[str], orders: dict[str, tuple]) -> None:
    """保持文本标识及原值；完整回读成功后才替换目标，失败保留原始下载。"""
    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_name(f".{target.stem}-{uuid4().hex}.xlsx")
    book = Workbook(write_only=True)
    sheet = book.create_sheet("出库单")
    try:
        sheet.append(headers)
        for values in orders.values():
            cells = []
            for value in values:
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cells.append(cell)
            sheet.append(cells)
        book.save(tmp)
        if read_orders(tmp) != (headers, orders):
            raise ValueError("订单文件回读校验失败")
        os.replace(tmp, target)
    finally:
        book.close()
        tmp.unlink(missing_ok=True)


def normalize_order_file(path: Path) -> int:
    headers, orders = read_orders(path, require_locations=True)
    write_orders(path, headers, orders)
    return len(orders)
